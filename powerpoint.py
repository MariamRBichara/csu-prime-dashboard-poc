from pathlib import Path
from datetime import datetime
import tempfile

import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


BASE_DIR = Path(__file__).parent
TEMPLATE_PATH = BASE_DIR / "templates" / "slide06_template.pptx"
OUTPUT_DIR = BASE_DIR / "output"


# Historical Jan-Jun trend retained from the approved reference.
# Source A values override these whenever available.
BASELINE_TREND = {
    1: 0.81,
    2: 0.82,
    3: 0.84,
    4: 0.85,
    5: 0.86,
    6: 0.87,
}


def _remove_shape(shape):
    sp = shape._element
    sp.getparent().remove(sp)


def _find_picture(slide, name):
    for shape in slide.shapes:
        if (
            shape.shape_type
            == MSO_SHAPE_TYPE.PICTURE
            and shape.name == name
        ):
            return shape

    raise ValueError(
        f"Required template picture '{name}' was not found."
    )


def _set_title(slide, reporting_period):
    wanted = (
        "CSU Prime Dashboard – "
        + reporting_period
    )

    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue

        if "CSU Prime Dashboard" in shape.text:
            frame = shape.text_frame

            if frame.paragraphs and frame.paragraphs[0].runs:
                first_run = frame.paragraphs[0].runs[0]
                first_run.text = wanted

                for run in frame.paragraphs[0].runs[1:]:
                    run.text = ""
            else:
                shape.text = wanted

            return

    raise ValueError(
        "Dashboard title was not found in the template."
    )


def _save_kpi_strip(metrics, output_path):
    k = metrics["kpis"]

    labels = [
        "HC tagged for ARVI",
        "Completion rate",
        "HC compliance",
        "ARVE",
        "ARVI",
        "Overbooking",
    ]

    values = [
        f"{int(k['hc_tagged'])}",
        f"{k['completion_rate']:.0%}",
        f"{k['hc_compliance']:.0%}",
        f"{k['arve']:.0%}",
        f"{k['arvi']:.0%}",
        f"{k['overbooking']:.0%}",
    ]

    value_colors = [
